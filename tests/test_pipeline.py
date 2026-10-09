"""End-to-end runs on the committed synthetic sample data."""

import openpyxl
from conftest import SAMPLE

from packing_spec.models import Status
from packing_spec.pipeline import run, summary_lines


def _by_sku(rows):
    return {r.sku: r for r in rows}


def test_demo_run_end_to_end(sample_config):
    result = run(sample_config)
    case = _by_sku(result.case_rows)
    dims = _by_sku(result.dimension_rows)

    assert result.instruction_order[0].startswith("1. PO 1042")
    assert case["FBA-HX-1042-P5"].instructions == 40  # newer PO overrides older 36
    assert case["MF-HX-1042-P20"].status is Status.UPDATE
    assert case["FBA-CC-2210-P10"].status is Status.REVIEW
    assert case["MF-TW-5150-P50"].status is Status.NOT_SHIPPED
    # Carton size missing in the newest PO, filled from the older one.
    assert dims["FBA-CC-2210-P20"].instructions_source == "PO 1041"
    assert dims["FBA-GR-4400-P6"].status is Status.OK and "CENTIMETRES" in dims[
        "FBA-GR-4400-P6"].note
    assert any("FBA-PT-7720-P12" in w for w in result.warnings)
    assert not result.missing_instructions


def test_merged_sku_cell_and_na_cell_rows_are_not_lost(sample_config):
    case = _by_sku(run(sample_config).case_rows)
    merged = case["FBA-MB-3305-P25"]
    assert len(merged.pl_rows) == 2 and merged.status is Status.REVIEW
    # "#N/A" in the internal-SKU column resolved via parent + pack, ERP breaks the tie.
    assert case["FBA-SL-0815-P4"].packing_list == 20


def test_outputs_have_live_formulas_and_enriched_instructions(sample_config):
    result = run(sample_config)
    ws = openpyxl.load_workbook(result.case_qty_path)["Case Qty"]
    assert ws["I1"].value == "Correct case qty"
    assert ws["M2"].value.startswith("=IF(")
    [enriched] = result.extra_outputs
    assert "WITH INTERNAL SKU" in enriched.name


def test_missing_instructions_asks_the_operator(sample_config, tmp_path):
    sample_config.input_dir = None
    sample_config.erp = SAMPLE / "input" / "erp_product_export_2026-09-18.xlsx"
    sample_config.packing_list = SAMPLE / "input" / "Northwind packing list PO 1041-1042 " \
                                                    "2026-09-20.xlsx"
    result = run(sample_config)
    assert result.missing_instructions
    assert "ASK THE OPERATOR" in "\n".join(summary_lines(result))
    assert all(r.instructions is None for r in result.case_rows)


def test_explicit_instruction_order_overrides_inference(sample_config):
    files = sorted((SAMPLE / "input").glob("Packing Instructions*.xlsx"))  # PO 1041 first
    sample_config.input_dir = None
    sample_config.erp = SAMPLE / "input" / "erp_product_export_2026-09-18.xlsx"
    sample_config.packing_list = next((SAMPLE / "input").glob("Northwind packing list*"))
    sample_config.instructions = files
    case = _by_sku(run(sample_config).case_rows)
    assert case["FBA-HX-1042-P5"].instructions == 36
    assert case["FBA-HX-1042-P5"].instructions_source == "PO 1041"
