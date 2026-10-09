"""Input discovery: decide which workbook is which by CONTENT, never by file name.

File names change with every shipment and PO, and operators drop extra files in the
input folder. Recognized roles:

    ERP export              a tab with ProductID + QtyPerCase              (required)
    packing instructions    a tab named "Box Information"                  (0..n, one per PO)
    vendor packing list     an item table: item/SKU + quantity             (needed to compare)

Skipped automatically: reports produced by this tool (and the instructions copy it
writes), and outputs of the sibling tools in the same workflow (invoice reconciliation,
unpacking list), which also look like item tables and would otherwise be mistaken for
the vendor packing list.

Check order matters because the heuristics overlap: a tab NAMED "Box Information" is
unambiguous, so it is tested before the item-table rule; "any tab with a bags/carton
column" is accepted as instructions only last, since the packing list matches it too.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .instructions import detect_instructions_sheet
from .packing_list import pick_packing_list_sheet
from .po_tags import order_by_priority, source_label
from .report import ADDED_SKU_HEADER
from .sheet_utils import find_header_row, normalize_header, row_values

log = logging.getLogger(__name__)

OWN_OUTPUT_HEADERS = {"correctcaseqty", "correctcasedimensionsmm",
                      normalize_header(ADDED_SKU_HEADER)}


@dataclass
class DiscoveredInputs:
    erp: Path
    packing_list: Path | None = None
    instructions: list[Path] = field(default_factory=list)
    sheets: dict[Path, str] = field(default_factory=dict)
    ignored: list[str] = field(default_factory=list)


def _header_sets(wb: openpyxl.Workbook) -> dict[str, set[str]]:
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = {find_header_row(ws), find_header_row(ws, max_scan=45, min_text_cells=4)}
        out[name] = {normalize_header(h) for r in rows for h in row_values(ws, r) if h}
    return out


def _is_own_output(headers: dict[str, set[str]]) -> bool:
    return any(h & OWN_OUTPUT_HEADERS for h in headers.values())


def _is_sibling_output(headers: dict[str, set[str]]) -> bool:
    for h in headers.values():
        if "priceinvoice" in h and any(x.startswith("qtypacking") for x in h):
            return True  # invoice-vs-packing reconciliation report
        if {"sku", "pcsbag", "ctns"} <= h and ({"bagqty", "bagsqty"} & h):
            return True  # unpacking list for the warehouse
    return False


def _erp_sheet(headers: dict[str, set[str]]) -> str | None:
    for name, h in headers.items():
        if h & {"productid", "internalsku"} and any("qtypercase" in x for x in h):
            return name
    return None


def classify_inputs(input_dir: Path) -> DiscoveredInputs:
    files = sorted(f for f in Path(input_dir).glob("*.xlsx") if not f.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No .xlsx files in {input_dir}.")

    erp: Path | None = None
    found = DiscoveredInputs(erp=Path())
    for path in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        headers = _header_sets(wb)
        if _is_own_output(headers):
            found.ignored.append(f"{path.name} (a report produced by this tool)")
            continue
        if _is_sibling_output(headers):
            found.ignored.append(f"{path.name} (output of a sibling tool)")
            continue
        if sheet := _erp_sheet(headers):
            erp = _claim(erp, path, "ERP export")
            found.sheets[path] = sheet
        elif sheet := detect_instructions_sheet(wb, by_name_only=True):
            found.instructions.append(path)
            found.sheets[path] = sheet
        elif sheet := pick_packing_list_sheet(wb):
            found.packing_list = _claim(found.packing_list, path, "vendor packing list")
            found.sheets[path] = sheet
        elif sheet := detect_instructions_sheet(wb):
            found.instructions.append(path)
            found.sheets[path] = sheet
        else:
            found.ignored.append(f"{path.name} (content not recognized; tabs {wb.sheetnames})")

    if erp is None:
        raise FileNotFoundError(
            f"No ERP product export in {input_dir} (looking for a tab with ProductID + "
            f"QtyPerCase). It is the only required input.")
    found.erp = erp
    found.instructions = order_by_priority(found.instructions)
    log_discovery(found)
    return found


def log_discovery(found: DiscoveredInputs) -> None:
    log.info("  ERP export:    %s", found.erp.name)
    log.info("  packing list:  %s", found.packing_list.name if found.packing_list else "MISSING")
    if not found.instructions:
        log.info("  instructions:  MISSING")
    for i, path in enumerate(found.instructions, start=1):
        rank = "WINS" if i == 1 else f"fills gaps (#{i})"
        log.info("  instructions [%s] -> %s: %s", source_label(path), rank, path.name)
    for line in found.ignored:
        log.info("  ignored: %s", line)


def _claim(current: Path | None, path: Path, role: str) -> Path:
    if current is not None:
        raise ValueError(f"Two files look like the {role}: '{current.name}' and '{path.name}'. "
                         f"Keep only one in the input folder.")
    return path
