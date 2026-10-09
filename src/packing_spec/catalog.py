"""SKU catalog: the bridge between the three SKU levels.

    parent SKU        short style code, no inventory          HX-1042
      internal SKU    channel prefix + pack size              FBA-HX-1042-P5   <- ERP ProductID
        marketplace   opaque, no structure                    QK7-2LM-9TZ

Vendors only ever see the parent and the marketplace SKU (it is printed on the carton
label). The internal SKU never leaves the company, yet it is the ERP's key, so every
file has to be translated to that level before anything can be compared.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from .models import Number
from .sheet_utils import find_column, normalize_header, require_column

log = logging.getLogger(__name__)


def pack_key(value: Any) -> int | str:
    try:
        return int(round(float(str(value).strip())))
    except (TypeError, ValueError):
        return str(value).strip()


@dataclass
class SkuCatalog:
    by_marketplace: dict[str, str] = field(default_factory=dict)
    by_parent_pack: dict[tuple[str, int | str], list[str]] = field(default_factory=dict)
    names: dict[str, str] = field(default_factory=dict)
    pack_sizes: dict[str, Number] = field(default_factory=dict)
    parents: dict[str, str] = field(default_factory=dict)

    def internal_from_marketplace(self, marketplace_sku: Any) -> str | None:
        if marketplace_sku is None:
            return None
        return self.by_marketplace.get(str(marketplace_sku).strip())

    def candidates(self, parent: Any, pack_size: Any) -> list[str]:
        """Every internal SKU for (parent, pieces per bag).

        Usually two (the FBA- and the MF- variant). The caller breaks the tie using the
        ERP export; this method never picks one arbitrarily.
        """
        if parent is None or pack_size is None:
            return []
        return list(self.by_parent_pack.get((str(parent).strip(), pack_key(pack_size)), []))


def load_catalog(path: Path) -> SkuCatalog:
    """Read the catalog sheet in read-only streaming mode.

    The real catalog workbook is large and multi-tab; only five columns of one tab are
    needed, and read-only mode avoids parsing everything else.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = _pick_sheet(wb)
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows))
        c_sku = require_column(headers, lambda h: h in ("internalsku", "sku"), "Internal SKU")
        c_parent = find_column(headers, lambda h: h in ("parentsku", "parent"))
        c_name = find_column(headers, lambda h: h in ("product", "productname"))
        c_pack = find_column(headers, lambda h: h == "packsize")
        c_mkt = find_column(headers, lambda h: h in ("marketplacesku", "amazonsku"))

        catalog = SkuCatalog()
        for row in rows:
            def cell(col: int | None, row: tuple = row) -> str | None:
                if not col or len(row) < col or row[col - 1] is None:
                    return None
                return str(row[col - 1]).strip() or None

            sku = cell(c_sku)
            if not sku:
                continue
            parent, pack = cell(c_parent), cell(c_pack)
            if parent:
                catalog.parents.setdefault(sku, parent)
            if name := cell(c_name):
                catalog.names.setdefault(sku, name)
            if pack:
                catalog.pack_sizes.setdefault(sku, pack_key(pack))
                if parent:
                    catalog.by_parent_pack.setdefault((parent, pack_key(pack)), []).append(sku)
            if mkt := cell(c_mkt):
                catalog.by_marketplace.setdefault(mkt, sku)
    finally:
        wb.close()
    log.info("  catalog: %d SKUs with pack size, %d marketplace SKUs mapped",
             len(catalog.pack_sizes), len(catalog.by_marketplace))
    return catalog


def _pick_sheet(wb: Any) -> Any:
    for name in wb.sheetnames:
        first = next(wb[name].iter_rows(max_row=1, values_only=True), ())
        if any(normalize_header(h) in ("internalsku", "sku") for h in first):
            return wb[name]
    return wb[wb.sheetnames[0]]
