"""Vendor packing list: what the vendor SAYS they packed.

Reads the item table, verifies the carton unit, maps every row to an internal SKU and
aggregates rows per SKU.

Internal SKU resolution, in order:

1. A column whose *values* exist as ERP product IDs (often it has no header at all,
   while the column named "FBA SKU" actually holds marketplace SKUs).
2. The marketplace-SKU column, translated through the catalog.
3. Parent + pieces per bag through the catalog. This usually yields two candidates
   (FBA- and MF- variants); the ERP export breaks the tie. If it cannot, the row is
   flagged with both names instead of being assigned arbitrarily.

Rows left blank or holding "#N/A" in the internal-SKU column fall through to step 3.

Several rows per SKU is normal when they are different SKUs of one parent; two rows
resolving to the SAME SKU with different bags/carton is a conflict: the value backed by
more cartons is shown, and the row is forced to manual review with both values listed.
"""

from __future__ import annotations

import logging
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .catalog import SkuCatalog
from .dims import extract_numbers, find_lwh_columns, fmt_box, resolve_metric_unit
from .models import Box, Number, PackingLine, ShippedSku
from .sheet_utils import (
    cell_through_merge,
    clean_sku,
    find_column,
    find_header_row,
    find_sku_column_by_content,
    merged_anchor_map,
    normalize_header,
    row_values,
    to_number,
)

log = logging.getLogger(__name__)


def is_parent_header(h: str) -> bool:
    return h.startswith("itemno") or h in ("parent", "parentsku", "style", "styleno")


def is_quantity_header(h: str) -> bool:
    return h.startswith("quant") or h in ("totalqty", "totalpieces", "qtypcs")


def is_marketplace_header(h: str) -> bool:
    return h in ("fbasku", "marketplacesku", "amazonsku")


@dataclass
class PackingListResult:
    lines: list[PackingLine]
    sku_method: str
    warnings: list[str] = field(default_factory=list)


def pick_packing_list_sheet(wb: openpyxl.Workbook) -> str | None:
    """The item table tab. A tab literally named 'Packing List' wins over other matches."""
    candidates = []
    for name in wb.sheetnames:
        headers = [normalize_header(h) for h in row_values(wb[name], find_header_row(wb[name]))]
        if any(is_parent_header(h) or h == "sku" for h in headers) and any(
            is_quantity_header(h) for h in headers
        ):
            candidates.append(name)
    named = [n for n in candidates if "packinglist" in normalize_header(n)]
    return (named or candidates or [None])[0]


def read_packing_list(
    path: Path, known_skus: set[str], catalog: SkuCatalog, sheet: str | None = None
) -> PackingListResult:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = sheet or pick_packing_list_sheet(wb)
    if sheet is None:
        raise ValueError(f"'{path.name}' has no item table (needs an item/SKU and a quantity "
                         f"column). Pass the sheet explicitly with --pl-sheet.")
    ws = wb[sheet]
    header_row = find_header_row(ws)
    headers = row_values(ws, header_row)

    cols = {
        "parent": find_column(headers, is_parent_header),
        "marketplace": find_column(headers, is_marketplace_header),
        "pcs": find_column(headers, lambda h: h in ("pcsbag", "pcsperbag", "piecesperbag")),
        "bags": find_column(headers, lambda h: ("bag" in h or "box" in h)
                            and ("ctn" in h or "carton" in h) and not h.startswith("pcs")),
        "ctns": find_column(headers, lambda h: h in ("ctns", "cartons", "noofctns", "totalctns")),
        "qty": find_column(headers, is_quantity_header),
        "size": find_column(headers, lambda h: h in ("cartonsize", "ctnsize", "sizectn")),
        "cbm": find_column(headers, lambda h: h.startswith("meas") or "cbm" in h),
    }
    lwh = next((t for t in find_lwh_columns(headers) if t.unit != "in"), None)
    sku_col, ratio = find_sku_column_by_content(ws, header_row, known_skus)

    identity_cols = [c for c in (cols["parent"], cols["marketplace"], cols["size"], sku_col) if c]
    anchors = merged_anchor_map(ws, identity_cols)

    lines: list[PackingLine] = []
    mislabelled = 0
    for r in range(header_row + 1, ws.max_row + 1):
        parent = clean_sku(cell_through_merge(ws, r, cols["parent"], anchors))
        internal = clean_sku(cell_through_merge(ws, r, sku_col, anchors))
        raw_internal = cell_through_merge(ws, r, sku_col, anchors)
        marketplace = clean_sku(cell_through_merge(ws, r, cols["marketplace"], anchors))
        if not (parent or raw_internal or marketplace):
            continue  # totals row, blank spacer or footnote
        bags = to_number(ws.cell(row=r, column=cols["bags"]).value) if cols["bags"] else None
        cartons = to_number(ws.cell(row=r, column=cols["ctns"]).value) if cols["ctns"] else None
        box, unit = _row_box(ws, r, cols, lwh, anchors, cartons)
        if unit == "mm" and lwh and lwh.unit == "cm":
            mislabelled += 1
        lines.append(PackingLine(
            row=r,
            parent=parent,
            pcs_per_bag=to_number(ws.cell(row=r, column=cols["pcs"]).value) if cols["pcs"]
            else None,
            bags_per_carton=bags,
            cartons=cartons,
            quantity=to_number(ws.cell(row=r, column=cols["qty"]).value) if cols["qty"]
            else None,
            box_mm=box,
            marketplace_sku=marketplace,
            internal_sku=internal,
        ))

    result = PackingListResult(lines, "")
    if mislabelled:
        result.warnings.append(
            f"Packing list dimension columns say (cm) but {mislabelled} row(s) hold millimetres "
            f"(verified against the declared CBM); read as mm.")
    if sku_col is not None:
        result.sku_method = f"column {sku_col} by content ({ratio:.0%} of values are ERP SKUs)"
    elif cols["marketplace"] and _translate_marketplace(lines, catalog) >= 0.6 * len(lines):
        result.sku_method = f"column {cols['marketplace']} (marketplace SKU) via catalog"
    else:
        for line in lines:
            line.internal_sku = None
        result.sku_method = "parent + pieces per bag via catalog"
    _fill_gaps(lines, catalog, known_skus, result.warnings)
    log.info("  packing list: sheet '%s', %d rows; internal SKU from %s",
             sheet, len(lines), result.sku_method)
    return result


def _row_box(ws: Worksheet, r: int, cols: dict, lwh, anchors: dict,
             cartons: Number | None) -> tuple[Box | None, str | None]:
    if lwh:
        nums = [to_number(ws.cell(row=r, column=c).value) for c in lwh.cols]
    elif cols["size"]:
        nums = extract_numbers(cell_through_merge(ws, r, cols["size"], anchors))
    else:
        return None, None
    if len(nums) < 3 or not all(isinstance(v, (int, float)) and v > 0 for v in nums[:3]):
        return None, None
    cbm = to_number(ws.cell(row=r, column=cols["cbm"]).value) if cols["cbm"] else None
    return resolve_metric_unit([float(v) for v in nums], cartons=cartons, declared_cbm=cbm)


def _translate_marketplace(lines: list[PackingLine], catalog: SkuCatalog) -> int:
    for line in lines:
        line.internal_sku = catalog.internal_from_marketplace(line.marketplace_sku)
    return sum(1 for line in lines if line.internal_sku)


def _fill_gaps(lines: list[PackingLine], catalog: SkuCatalog, known: set[str],
               warnings: list[str]) -> None:
    """Resolve remaining rows by parent + pieces per bag, using the ERP as tie-breaker."""
    for line in (ln for ln in lines if not ln.internal_sku):
        candidates = catalog.candidates(line.parent, line.pcs_per_bag)
        in_erp = [c for c in candidates if c in known]
        if len(in_erp) == 1:
            line.internal_sku = in_erp[0]
        elif len(candidates) == 1:
            line.internal_sku = candidates[0]
        else:
            detail = (f"candidates {', '.join(candidates)} are "
                      f"{'all' if in_erp else 'not'} in the ERP export"
                      if candidates else "the catalog has no SKU for that parent + pack size")
            warnings.append(f"Packing list row {line.row}: could not resolve the internal SKU "
                            f"(parent '{line.parent}', {line.pcs_per_bag} pcs/bag); {detail}. "
                            f"Row excluded from the comparison.")


def aggregate_by_sku(lines: list[PackingLine]) -> dict[str, ShippedSku]:
    """Group physical rows by internal SKU, weighting conflicting values by cartons."""
    groups: dict[str, list[PackingLine]] = {}
    for line in lines:
        if line.internal_sku:
            groups.setdefault(line.internal_sku, []).append(line)

    out = {}
    for sku, rows in groups.items():
        case_votes: Counter = Counter()
        box_votes: Counter = Counter()
        alerts = []
        for line in rows:
            weight = line.cartons if isinstance(line.cartons, (int, float)) and line.cartons > 0 \
                else 1
            if line.bags_per_carton is not None:
                case_votes[line.bags_per_carton] += weight
            if line.box_mm:
                box_votes[tuple(round(v, 1) for v in line.box_mm)] += weight
            expected = _product(line.pcs_per_bag, line.bags_per_carton, line.cartons)
            if expected is not None and line.quantity and abs(expected - line.quantity) > 0.5:
                alerts.append(f"row {line.row}: pcs/bag x bags/ctn x ctns = {expected:g} but "
                              f"quantity says {line.quantity:g}")

        case_conflict = box_conflict = None
        if len(case_votes) > 1:
            detail = " and ".join(f"{v:g} bags/carton in {w:g} carton(s)"
                                  for v, w in case_votes.most_common())
            case_conflict = (f"vendor shipped this SKU with MORE THAN ONE carton packing "
                             f"({detail}); the majority value is shown")
        if len(box_votes) > 1:
            detail = " and ".join(f"{fmt_box(v)} in {w:g} carton(s)"
                                  for v, w in box_votes.most_common())
            box_conflict = (f"vendor shipped this SKU in DIFFERENT carton sizes ({detail}); "
                            f"the majority size is shown")

        out[sku] = ShippedSku(
            sku=sku,
            case_qty=case_votes.most_common(1)[0][0] if case_votes else None,
            case_qty_values=[v for v, _ in case_votes.most_common()],
            box_mm=box_votes.most_common(1)[0][0] if box_votes else None,
            cartons=sum(ln.cartons for ln in rows if isinstance(ln.cartons, (int, float))),
            quantity=sum(ln.quantity for ln in rows if isinstance(ln.quantity, (int, float))),
            parent=rows[0].parent,
            rows=[ln.row for ln in rows],
            case_conflict=case_conflict,
            box_conflict=box_conflict,
            alerts=alerts,
        )
    log.info("  packing list: %d rows -> %d internal SKUs", len(lines), len(out))
    return out


def _product(*values: Number | None) -> float | None:
    if not all(isinstance(v, (int, float)) for v in values):
        return None
    result = 1.0
    for v in values:
        result *= v
    return result
