"""Packing instructions: what WE asked the vendor for when the PO was placed.

The workbook has a "Box Information" tab with two facts per SKU: bags per carton and
the carton size. Quirks handled here:

* The header is not on row 1: half a page of prose instructions for the vendor sits on
  top, one text cell per row, so the header is searched for up to row 45.
* The file is written for the vendor, so it carries the parent and marketplace SKUs but
  usually not the internal SKU. It is located by content when present, otherwise
  translated from the marketplace SKU via the catalog (and a copy of the file with the
  column added is written to the output folder).
* Dimension columns titled "(cm)" may hold millimetres; the inch columns of the same
  row decide.
* A container can carry goods from several POs, each with its own instructions file.
  ``merge_instructions`` resolves them field by field: the newest PO wins, older files
  only fill gaps, and the source of every value is kept for the audit column.
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .catalog import SkuCatalog
from .dims import (
    MM_PER_INCH,
    LwhColumns,
    extract_numbers,
    resolve_metric_unit,
    split_metric_and_inch,
)
from .models import Box, InstructionEntry, InstructionsFile, MergedInstruction
from .sheet_utils import (
    clean_sku,
    find_column,
    find_header_row,
    find_sku_column_by_content,
    normalize_header,
    row_values,
    to_number,
)

log = logging.getLogger(__name__)

SHEET_NAME_TOKENS = ("boxinformation", "boxinfo")
MAX_HEADER_SCAN = 45


def is_bags_per_carton(h: str) -> bool:
    """'Bags (Boxes) per Carton', 'BAGS (BOXES)/CASE', 'bags/ctn', and the 'caton' typo."""
    if "qtypercase" in h:
        return True
    has_unit = "bag" in h or "box" in h
    has_carton = any(t in h for t in ("carton", "caton", "case", "ctn"))
    return has_unit and has_carton


def find_instructions_header_row(ws: Worksheet) -> int | None:
    for r in range(1, min(MAX_HEADER_SCAN, ws.max_row) + 1):
        headers = row_values(ws, r)
        texts = [h for h in headers if isinstance(h, str) and h.strip()]
        if len(texts) >= 4 and find_column(headers, is_bags_per_carton) is not None:
            return r
    return None


def detect_instructions_sheet(wb: Workbook, by_name_only: bool = False) -> str | None:
    """The "Box Information" tab, or (unless `by_name_only`) any tab with a bags/carton column.

    File discovery uses `by_name_only=True` first because the vendor packing list also
    has a bags-per-carton column; mistaking one file for the other silently corrupts the
    whole run.
    """
    for name in wb.sheetnames:
        if any(t in normalize_header(name) for t in SHEET_NAME_TOKENS):
            return name
    if by_name_only:
        return None
    for name in wb.sheetnames:
        if find_instructions_header_row(wb[name]) is not None:
            return name
    return None


def load_instructions(
    path: Path,
    label: str,
    known_skus: set[str],
    catalog: SkuCatalog,
    sheet: str | None = None,
) -> InstructionsFile | None:
    """Parse one instructions workbook. Returns None (with a logged reason) if unusable."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = sheet or detect_instructions_sheet(wb)
    if sheet is None:
        log.warning("  '%s' has no Box Information tab; skipped", path.name)
        return None
    ws = wb[sheet]
    header_row = find_instructions_header_row(ws) or find_header_row(ws)
    headers = row_values(ws, header_row)
    qty_col = find_column(headers, is_bags_per_carton)
    if qty_col is None:
        log.warning("  '%s': no bags-per-carton column in tab '%s'; skipped", path.name, sheet)
        return None
    metric, inch = split_metric_and_inch(headers)

    # Lower ratio floor than the packing list: this file lists every SKU of the PO,
    # not just the ones on this container.
    sku_col, ratio = find_sku_column_by_content(ws, header_row, known_skus, min_ratio=0.25)
    internal_sku_added = sku_col is None
    if sku_col is not None:
        source_col = sku_col
        log.info("  %s: internal SKU found by content in column %d (%.0f%% known to ERP)",
                 label, sku_col, ratio * 100)
    else:
        source_col = find_column(
            headers, lambda h: h in ("fbasku", "marketplacesku") or "amazon" in h
        )
        if source_col is None:
            log.warning("  '%s': neither an internal SKU nor a marketplace SKU column; skipped",
                        path.name)
            return None
        log.info("  %s: no internal SKU column; translating marketplace SKUs via catalog", label)

    entries: dict[str, InstructionEntry] = {}
    row_skus: list[tuple[int, str | None]] = []
    unresolved: list[int] = []
    duplicates: list[str] = []
    mislabelled = 0
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=source_col).value
        sku = clean_sku(raw) if sku_col else catalog.internal_from_marketplace(clean_sku(raw))
        row_skus.append((r, sku))
        if sku is None:
            # A totals row has numbers but no SKU; only a row WITH a SKU that failed to
            # translate is a real problem.
            if clean_sku(raw):
                unresolved.append(r)
            continue
        bags = to_number(ws.cell(row=r, column=qty_col).value)
        box, unit = _row_box(ws, r, metric, inch)
        if bags is None and box is None:
            continue
        if unit == "mm" and metric and metric.unit == "cm":
            mislabelled += 1
        if sku in entries:
            if bags is not None and entries[sku].bags_per_carton not in (None, bags):
                duplicates.append(f"{sku} ({entries[sku].bags_per_carton} vs {bags}, row {r})")
            continue
        entries[sku] = InstructionEntry(bags, box, r)

    warnings = []
    if mislabelled:
        warnings.append(f"'{path.name}': dimension columns say (cm) but {mislabelled} row(s) "
                        f"hold millimetres (verified against the inch columns); read as mm.")
    if unresolved:
        warnings.append(f"'{path.name}': {len(unresolved)} row(s) have a SKU that could not be "
                        f"translated to an internal SKU (rows {_short_list(unresolved)}).")
    if duplicates:
        warnings.append(f"'{path.name}': SKU listed twice with different bags/carton, first "
                        f"one kept: {', '.join(duplicates)}.")
    log.info("  %s: tab '%s', header row %d, %d SKUs", label, sheet, header_row, len(entries))
    return InstructionsFile(path, label, sheet, header_row, entries, internal_sku_added,
                            row_skus, warnings)


def _row_box(
    ws: Worksheet, r: int, metric: LwhColumns | None, inch: LwhColumns | None
) -> tuple[Box | None, str | None]:
    inches = (None, None, None)
    if inch:
        inches = tuple(to_number(ws.cell(row=r, column=c).value) for c in inch.cols)
    if metric:
        nums = [to_number(ws.cell(row=r, column=c).value) for c in metric.cols]
        if all(isinstance(v, (int, float)) and v > 0 for v in nums):
            return resolve_metric_unit(extract_numbers(" x ".join(map(str, nums))), inches)
    if inch and all(isinstance(v, (int, float)) and v > 0 for v in inches):
        return tuple(v * MM_PER_INCH for v in inches), "in"
    return None, None


def merge_instructions(
    files: list[InstructionsFile],
) -> tuple[dict[str, MergedInstruction], dict[str, dict[str, int]]]:
    """Combine instruction files given in priority order (first = newest PO).

    Resolution is per SKU AND per field: if the newest PO lists a SKU without a carton
    size, the size comes from the older PO and the provenance column says so.
    """
    merged: dict[str, MergedInstruction] = {}
    stats = {f.label: {"bags": 0, "box": 0} for f in files}
    for f in files:
        for sku, entry in f.entries.items():
            m = merged.setdefault(sku, MergedInstruction())
            if m.bags_per_carton is None and entry.bags_per_carton is not None:
                m.bags_per_carton, m.bags_source = entry.bags_per_carton, f.label
                stats[f.label]["bags"] += 1
            if m.box_mm is None and entry.box_mm is not None:
                m.box_mm, m.box_source = entry.box_mm, f.label
                stats[f.label]["box"] += 1
            if f.label not in m.sources:
                m.sources.append(f.label)
    return merged, stats


def describe_merge(files: list[InstructionsFile], stats: dict, merged: dict) -> list[str]:
    lines = []
    if len(files) > 1:
        lines.append("Priority: " + " > ".join(f.label for f in files)
                     + " (first wins; later files only fill gaps)")
    for f in files:
        s = stats[f.label]
        lines.append(f"  from {f.label}: {s['bags']} bags/carton values, {s['box']} carton sizes")
    lines.append(f"  combined: {len(merged)} SKUs")
    return lines


def _short_list(rows: list[int], limit: int = 15) -> str:
    return ", ".join(map(str, rows[:limit])) + ("..." if len(rows) > limit else "")
