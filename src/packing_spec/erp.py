"""ERP product export: what the system currently believes about each SKU's packing.

This file is the backbone of both reports: every output row is one SKU of this export,
in export order. Two numbers look alike but must never be confused:

* ``QtyPerCase``  - bags per carton the VENDOR should ship. Compared.
* ``FBACASEQTY``  - how WE re-pack for the marketplace. Shown as reference only.

The ERP stores the vendor carton twice, and the two copies drift apart: a free-text
``SUPPLIERCASEDIMENSIONS`` (mm or cm) and numeric ``CaseLength/Width/Height`` (inches).
``SUPPLIERCASEQTY`` is a deprecated custom column that reads "Ignore" on every row; it
is carried through verbatim and, being text, can never win a comparison.
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from .models import ErpProduct
from .sheet_utils import find_column, find_header_row, require_column, row_values, to_number

log = logging.getLogger(__name__)

SKU_HEADERS = ("productid", "sku", "internalsku")


def load_erp_export(path: Path, sheet: str | None = None) -> tuple[list[ErpProduct], list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    header_row = find_header_row(ws)
    headers = row_values(ws, header_row)

    c_sku = require_column(headers, lambda h: h in SKU_HEADERS, "ProductID")
    c_qpc = require_column(headers, lambda h: "qtypercase" in h, "QtyPerCase")
    optional = {
        "ProductName": find_column(headers, lambda h: h in ("productname", "name", "product")),
        "FBACASEQTY": find_column(headers, lambda h: "fbacase" in h),
        "SUPPLIERCASEQTY": find_column(headers, lambda h: "suppliercaseqty" in h),
        "SUPPLIERCASEDIMENSIONS": find_column(headers, lambda h: "suppliercasedimension" in h),
        "CaseLength": find_column(headers, lambda h: h == "caselength"),
        "CaseWidth": find_column(headers, lambda h: h == "casewidth"),
        "CaseHeight": find_column(headers, lambda h: h == "caseheight"),
    }

    def val(r: int, col: int | None) -> object:
        return ws.cell(row=r, column=col).value if col else None

    products = []
    for r in range(header_row + 1, ws.max_row + 1):
        sku = val(r, c_sku)
        if sku is None or not str(sku).strip():
            continue
        name = val(r, optional["ProductName"])
        products.append(ErpProduct(
            sku=str(sku).strip(),
            name=str(name).strip() if name else None,
            fba_case_qty=to_number(val(r, optional["FBACASEQTY"])),
            qty_per_case=to_number(val(r, c_qpc)),
            supplier_case_qty=val(r, optional["SUPPLIERCASEQTY"]),
            dims_text=val(r, optional["SUPPLIERCASEDIMENSIONS"]),
            case_inches=tuple(
                to_number(val(r, optional[k])) for k in ("CaseLength", "CaseWidth", "CaseHeight")
            ),
            row=r,
        ))

    missing = [label for label, col in optional.items() if col is None and label != "ProductName"]
    warnings = [f"ERP export has no '{label}' column; that part of the report stays empty."
                for label in missing]
    log.info("  ERP export: sheet '%s', header row %d, %d SKUs", ws.title, header_row,
             len(products))
    return products, warnings
