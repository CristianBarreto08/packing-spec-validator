"""Excel outputs: one workbook per question, plus a 'Run info' tab for traceability.

Case qty keeps the column order of the spreadsheet template buyers already used, with
"Correct case qty" in column I (the column they copy into the ERP); new columns are
appended after N rather than inserted.

The audit columns (diff, J..N match checks, inch conversions) are live formulas on the
same sheet, so a reviewer can see and re-run the rule. The decision columns
(Correct, Status) are values: they encode rules the template never had (multi-packing
conflicts, vendor-vs-instructions), and no formula points at another workbook, so the
files keep working after being emailed or moved.

Colors: green = OK, amber = value known but an ERP field must be updated, red = review
manually, grey = not on this packing list. Every flagged row carries a cell comment with
the full reason and the packing-list rows it came from.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .dims import fmt_box
from .models import CaseQtyRow, DimensionRow, InstructionsFile, Status

log = logging.getLogger(__name__)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


HEADER_FILL = _fill("FF1F3A5F")
GROUP_FILLS = {"ours": _fill("FF2F6F9F"), "vendor": _fill("FF8A6D00"), "correct": _fill("FF2E7D32")}
STATUS_FILLS = {
    Status.OK: _fill("FFC6EFCE"),
    Status.UPDATE: _fill("FFFFE699"),
    Status.REVIEW: _fill("FFFFC7CE"),
    Status.NOT_SHIPPED: _fill("FFEDEDED"),
}
CORRECT_FILL = _fill("FFE2F0D9")
THIN = Border(bottom=Side(style="thin", color="FFBFBFBF"))


@dataclass(frozen=True)
class Col:
    header: str
    width: int = 12
    group: str | None = None
    fmt: str | None = None


CASE_QTY_COLUMNS = [
    Col("Product ID", 20), Col("Product name", 28), Col("Pack size", 8, fmt="0"),
    Col("QtyPerCase (ERP)", 11, fmt="0"),
    Col("Supplier case qty (custom, deprecated)", 14),
    Col("Instructions we sent", 12, "ours", "0"),
    Col("PL case qty", 10, "vendor", "0"),
    Col("Diff PL - instructions", 11, "vendor", "0"),
    Col("Correct case qty", 11, "correct", "0"),
    Col("All four match", 10, fmt="0"),
    Col("Custom = Instructions = PL", 12, fmt="0"),
    Col("Custom = PL", 10, fmt="0"),
    Col("Instructions = PL", 11, fmt="0"),
    Col("ERP = PL", 10, fmt="0"),
    Col("FBA case qty (reference only)", 12, fmt="0"),
    Col("Instructions source (PO)", 13, "ours"),
    Col("Vendor vs instructions", 24, "ours"),
    Col("Status", 13), Col("What to do", 70),
]

DIMENSION_COLUMNS = [
    Col("Product ID", 20), Col("Product name", 28), Col("Pack size", 8, fmt="0"),
    Col("FBA case qty (reference only)", 12, fmt="0"),
    Col("ERP size text (raw)", 24),
    Col("ERP text L (mm)", 9, fmt="0"), Col("ERP text W (mm)", 9, fmt="0"),
    Col("ERP text H (mm)", 9, fmt="0"),
    Col("ERP CaseLength (in)", 10, fmt="0.00"), Col("ERP CaseWidth (in)", 10, fmt="0.00"),
    Col("ERP CaseHeight (in)", 10, fmt="0.00"),
    Col("Instructions we sent (mm)", 14, "ours"),
    Col("PL carton (mm)", 14, "vendor"),
    Col("PL L (mm)", 8, "vendor", "0"), Col("PL W (mm)", 8, "vendor", "0"),
    Col("PL H (mm)", 8, "vendor", "0"),
    Col("PL L (in)", 8, "vendor", "0.00"), Col("PL W (in)", 8, "vendor", "0.00"),
    Col("PL H (in)", 8, "vendor", "0.00"),
    Col("Correct case dimensions (mm)", 15, "correct"),
    Col("ERP text vs PL", 16), Col("ERP inches vs PL", 16),
    Col("Instructions vs PL", 16, "ours"), Col("Instructions source (PO)", 13, "ours"),
    Col("Status", 13), Col("What to do", 70),
]


def _write_table(ws: Worksheet, columns: list[Col]) -> None:
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.header)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = GROUP_FILLS.get(col.group, HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"


def _write_row(ws: Worksheet, r: int, columns: list[Col], values: list[Any],
               status: Status, note: str, pl_rows: list[int], correct_col: int) -> None:
    status_fill = STATUS_FILLS[status]
    for c, (col, value) in enumerate(zip(columns, values, strict=True), start=1):
        cell = ws.cell(row=r, column=c, value=value)
        cell.border = THIN
        if col.fmt:
            cell.number_format = col.fmt
        if status is not Status.OK:
            cell.fill = status_fill
    ws.cell(row=r, column=len(columns) - 1).fill = status_fill
    ws.cell(row=r, column=len(columns) - 1).font = Font(bold=True)
    ws.cell(row=r, column=len(columns)).alignment = Alignment(wrap_text=False)
    if status in (Status.OK, Status.UPDATE) and ws.cell(row=r, column=correct_col).value:
        ws.cell(row=r, column=correct_col).fill = CORRECT_FILL
        ws.cell(row=r, column=correct_col).font = Font(bold=True)
    if status is not Status.OK and note:
        where = f"\n\nPacking list rows: {', '.join(map(str, pl_rows))}" if pl_rows else ""
        ws.cell(row=r, column=len(columns)).comment = Comment(
            f"{status.value}: {note}{where}", "packing-spec", width=420, height=180
        )


def _finish(ws: Worksheet, columns: list[Col], last_row: int) -> None:
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(last_row, 2)}"


def _num_or_blank(ref_a: str, ref_b: str, expr: str, value: str) -> str:
    return f'=IF(AND(ISNUMBER({ref_a}),ISNUMBER({ref_b}),{expr}),{value},"")'


def case_qty_formulas(r: int) -> dict[int, str]:
    """Live formulas for H and J..N, mirroring `compare.resolve_case_qty`."""
    d, e, f, g = (f"{c}{r}" for c in "DEFG")
    nums = f"ISNUMBER({d}),ISNUMBER({e}),ISNUMBER({f}),ISNUMBER({g})"
    return {
        8: f'=IF(AND(ISNUMBER({f}),ISNUMBER({g})),{g}-{f},"")',
        10: f'=IF(AND({nums},{d}={e},{d}={f},{d}={g}),{d},"")',
        11: f'=IF(AND(ISNUMBER({e}),ISNUMBER({f}),ISNUMBER({g}),{e}={f},{e}={g}),{e},"")',
        12: _num_or_blank(e, g, f"{e}={g}", e),
        13: _num_or_blank(f, g, f"{f}={g}", f),
        14: _num_or_blank(d, g, f"{d}={g}", d),
    }


def write_case_qty_report(rows: list[CaseQtyRow], path: Path, run_info: list[str]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Qty"
    _write_table(ws, CASE_QTY_COLUMNS)
    for r, row in enumerate(rows, start=2):
        values: list[Any] = [
            row.sku, row.name, row.pack_size, row.erp_qty_per_case, row.supplier_case_qty,
            row.instructions, row.packing_list, None, row.correct,
            None, None, None, None, None,
            row.fba_case_qty, row.instructions_source, row.vendor_vs_instructions,
            row.status.value, row.note,
        ]
        for c, formula in case_qty_formulas(r).items():
            values[c - 1] = formula
        _write_row(ws, r, CASE_QTY_COLUMNS, values, row.status, row.note, row.pl_rows, 9)
    _finish(ws, CASE_QTY_COLUMNS, len(rows) + 1)
    _write_run_info(wb, run_info, [r.status for r in rows])
    wb.save(path)
    log.info("  wrote %s (%d SKUs)", path.name, len(rows))
    return path


def write_dimension_report(rows: list[DimensionRow], path: Path, run_info: list[str]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dimensions"
    _write_table(ws, DIMENSION_COLUMNS)
    for r, row in enumerate(rows, start=2):
        text_mm = row.erp_text_mm or (None, None, None)
        pl_mm = row.packing_list_mm or (None, None, None)
        inch_formulas = [
            f'=IF(ISNUMBER({c}{r}),ROUND({c}{r}/25.4,2),"")' if row.packing_list_mm else None
            for c in "NOP"
        ]
        values = [
            row.sku, row.name, row.pack_size, row.fba_case_qty, row.erp_text,
            *[_round(v) for v in text_mm], *[_round(v, 2) for v in row.erp_inches],
            fmt_box(row.instructions_mm), fmt_box(row.packing_list_mm),
            *[_round(v) for v in pl_mm], *inch_formulas,
            fmt_box(row.correct_mm), row.text_vs_pl, row.inches_vs_pl, row.instructions_vs_pl,
            row.instructions_source, row.status.value, row.note,
        ]
        _write_row(ws, r, DIMENSION_COLUMNS, values, row.status, row.note, row.pl_rows, 20)
    _finish(ws, DIMENSION_COLUMNS, len(rows) + 1)
    _write_run_info(wb, run_info, [r.status for r in rows])
    wb.save(path)
    log.info("  wrote %s (%d SKUs)", path.name, len(rows))
    return path


def _write_run_info(wb: openpyxl.Workbook, lines: list[str], statuses: list[Status]) -> None:
    ws = wb.create_sheet("Run info")
    ws.column_dimensions["A"].width = 120
    ws.cell(row=1, column=1, value="Status counts").font = Font(bold=True)
    r = 2
    for status in Status:
        cell = ws.cell(row=r, column=1, value=f"{status.value}: "
                       f"{sum(1 for s in statuses if s is status)}")
        cell.fill = STATUS_FILLS[status]
        r += 1
    ws.cell(row=r + 1, column=1, value="Inputs and warnings").font = Font(bold=True)
    for i, line in enumerate(lines, start=r + 2):
        ws.cell(row=i, column=1, value=line)


def write_instructions_with_internal_sku(info: InstructionsFile, out_dir: Path,
                                         name_tag: str) -> Path:
    """Copy of an instructions file with the internal-SKU column appended.

    Only produced when the file lacked that column and it had to be derived from the
    catalog, so the buyer can keep the enriched version for the next shipment.
    """
    wb = openpyxl.load_workbook(info.path)
    ws = wb[info.sheet]
    col = ws.max_column + 1
    header = ws.cell(row=info.header_row, column=col, value=ADDED_SKU_HEADER)
    header.font = Font(bold=True)
    for r, sku in info.row_skus:
        if sku:
            ws.cell(row=r, column=col, value=sku)
    ws.column_dimensions[get_column_letter(col)].width = 22
    path = out_dir / f"Packing Instructions {name_tag} - WITH INTERNAL SKU.xlsx"
    wb.save(path)
    log.info("  wrote %s", path.name)
    return path


ADDED_SKU_HEADER = "Internal SKU (added)"


def _round(value: Any, digits: int = 0) -> Any:
    if not isinstance(value, (int, float)):
        return None
    return int(round(value)) if digits == 0 else round(value, digits)
