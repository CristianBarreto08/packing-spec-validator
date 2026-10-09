"""Generate the synthetic sample inputs used by the demo and the README.

Deterministic (seeded). Every value is invented. The data deliberately reproduces the
situations the tool exists for:

* ERP QtyPerCase stale, exposed only by the instructions file      -> UPDATE
* vendor packed differently than instructed                        -> REVIEW
* one SKU shipped in two carton packings (merged SKU cell)         -> REVIEW
* "#N/A" in the internal-SKU column, recovered via parent + pack   -> resolved
* packing-list SKU missing from the ERP export                     -> warning
* ERP size text in cm, dates inside the text, inch fields stale or 0
* carton within tolerance, sides rotated, vendor used another carton
* two instruction files (PO 1041 older, PO 1042 newer), newest wins per field
* PO 1042 file has no internal-SKU column -> enriched copy written
* both instructions and packing list title dimension columns "(cm)" but hold mm
* decoy files in the input folder: an old report and a sibling tool's output

Run:  python examples/generate_sample_data.py
"""

from __future__ import annotations

import random
import string
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "sample_data"
INPUT, REFERENCE = OUT / "input", OUT / "reference"
MM_PER_INCH = 25.4
rng = random.Random(1042)


@dataclass
class Sku:
    sku: str
    parent: str
    name: str
    pack: int
    pcs_per_bag: int
    in_erp: bool = True
    erp_qty: int | None = None
    erp_text: str | None = None
    erp_inches: tuple | None = None  # None = derive from true box
    instr_1041: tuple | None = None  # (bags, box or None)
    instr_1042: tuple | None = None
    pl_rows: list = field(default_factory=list)  # [(bags, ctns, box)]
    na_in_pl: bool = False
    qty_typo: bool = False


def inches(box):
    return tuple(round(v / MM_PER_INCH, 4) for v in box)


B1, B2, B3 = (450, 320, 210), (620, 410, 380), (520, 360, 300)

SKUS = [
    Sku("FBA-HX-1042-P5", "HX-1042", "Hex Key Set 10pc, 5-pack", 5, 5, erp_qty=40,
        erp_text="450 x 320 x 210, Confirmed", instr_1041=(36, B1), instr_1042=(40, B1),
        pl_rows=[(40, 12, B1)]),
    Sku("MF-HX-1042-P20", "HX-1042", "Hex Key Set 10pc, 20-pack", 20, 20, erp_qty=12,
        erp_text="480 x 330 x 250 - Confirmed", instr_1042=(10, (480, 330, 250)),
        pl_rows=[(10, 6, (480, 330, 250))]),
    Sku("FBA-CC-2210-P10", "CC-2210", "Cable Clips 50ct, 10-pack", 10, 10, erp_qty=24,
        erp_text="400 x 300 x 280", instr_1042=(24, (400, 300, 280)),
        pl_rows=[(30, 8, (400, 300, 280))]),
    Sku("FBA-CC-2210-P20", "CC-2210", "Cable Clips 50ct, 20-pack", 20, 20, erp_qty=15,
        erp_text="410 x 310 x 300 Confirmed", instr_1041=(12, (410, 310, 300)),
        instr_1042=(15, None), pl_rows=[(15, 4, (410, 310, 300))]),
    Sku("FBA-MB-3305-P25", "MB-3305", "Mesh Bags Small, 25-pack", 25, 25, erp_qty=60,
        erp_text="560 x 350 x 360 - Confirmed", instr_1041=(60, (560, 350, 360)),
        pl_rows=[(60, 5, (560, 350, 360)), (48, 2, (560, 350, 360))]),
    Sku("FBA-SL-0815-P4", "SL-0815", "Silicone Lids 4pc", 4, 4, erp_qty=20,
        erp_text="380 x 380 x 300", instr_1042=(20, (380, 380, 300)),
        pl_rows=[(20, 10, (380, 380, 300))], na_in_pl=True),
    Sku("MF-SL-0815-P4", "SL-0815", "Silicone Lids 4pc", 4, 4, in_erp=False),
    Sku("FBA-PT-7720-P12", "PT-7720", "Plant Ties 12pk", 12, 12, in_erp=False,
        pl_rows=[(50, 3, (300, 250, 200))]),
    Sku("MF-TW-5150-P50", "TW-5150", "Twist Ties 50pc", 50, 50, erp_qty=100,
        erp_text="300 x 200 x 200"),
    Sku("FBA-GR-4400-P6", "GR-4400", "Grip Pads 6pc", 6, 6, erp_qty=36,
        erp_text="45 x 32 x 34 Accurate", erp_inches=inches((450, 320, 340)),
        instr_1042=(36, (450, 320, 340)), pl_rows=[(36, 7, (450, 320, 340))]),
    Sku("FBA-GR-4400-P12", "GR-4400", "Grip Pads 6pc, 12-pack", 12, 12, erp_qty=18,
        erp_text="520 x 360 x 300 Confirmed", instr_1042=(18, (500, 350, 300)),
        pl_rows=[(18, 5, B3)]),
    Sku("FBA-WB-6001-P2", "WB-6001", "Wire Baskets 2pc", 2, 2, erp_qty=8,
        erp_text="620 x 410 x 380 - Confirmed 08/07/2026", erp_inches=(7.0, 16.1, 15.0),
        instr_1041=(8, B2), pl_rows=[(8, 9, B2)]),
    Sku("FBA-WB-6001-P1", "WB-6001", "Wire Basket single", 1, 1, erp_qty=16,
        erp_text="620 x 410 x 380", erp_inches=(24.41, 16.14, 0),
        instr_1041=(16, B2), pl_rows=[(16, 4, B2)]),
    Sku("FBA-SC-1290-P8", "SC-1290", "Storage Cubes, 8-pack", 8, 8, erp_qty=40,
        erp_text="450 x 320 x 210", instr_1042=(40, B1), pl_rows=[(40, 6, (452, 321, 208))],
        qty_typo=True),
    Sku("FBA-SC-1290-P16", "SC-1290", "Storage Cubes, 16-pack", 16, 16, erp_qty=20,
        erp_text="320 x 450 x 210", erp_inches=inches((320, 450, 210)),
        instr_1042=(20, B1), pl_rows=[(20, 6, B1)]),
    Sku("FBA-LB-3030-P3", "LB-3030", "Label Holders 3pc", 3, 3, erp_qty=24,
        erp_text="400 x 300 x 300", erp_inches=(12.0, 10.0, 8.0),
        instr_1042=(24, (450, 320, 210)), pl_rows=[(24, 5, B1)]),
    Sku("FBA-HX-1042-P10", "HX-1042", "Hex Key Set 10pc, 10-pack", 10, 10, erp_qty=30,
        erp_text="-", erp_inches=(0, 0, 0), pl_rows=[(36, 3, (500, 340, 260))]),
    # Other SKUs on the POs that did not travel in this container (not in the ERP export).
    Sku("FBA-KN-9910-P2", "KN-9910", "Knob Set 2pc", 2, 2, in_erp=False,
        instr_1041=(48, (400, 300, 250))),
    Sku("MF-KN-9910-P2", "KN-9910", "Knob Set 2pc", 2, 2, in_erp=False,
        instr_1041=(48, (400, 300, 250))),
    Sku("FBA-HK-2020-P6", "HK-2020", "Wall Hooks 6pc", 6, 6, in_erp=False,
        instr_1041=(30, (350, 300, 250)), instr_1042=(30, (350, 300, 250))),
]


def marketplace_sku() -> str:
    chars = string.ascii_uppercase + string.digits
    return "-".join("".join(rng.choice(chars) for _ in range(3)) for _ in range(3))


MKT = {s.sku: marketplace_sku() for s in SKUS}


def _save(wb: openpyxl.Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"  wrote {path.relative_to(OUT.parent)}")


def write_catalog() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog"
    ws.append(["Internal SKU", "Parent SKU", "Product", "Pack size", "Marketplace SKU"])
    for s in SKUS:
        ws.append([s.sku, s.parent, s.name, s.pack, MKT[s.sku]])
    _save(wb, REFERENCE / "sku_catalog.xlsx")


def write_erp_export() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ProductID", "ProductName", "FBACASEQTY", "QtyPerCase", "SUPPLIERCASEQTY",
               "SUPPLIERCASEDIMENSIONS", "CaseLength", "CaseWidth", "CaseHeight",
               "PACKINGINSTRUCTIONS", "PACKINGINSTRUCTIONS2"])
    for s in (s for s in SKUS if s.in_erp):
        true_box = s.pl_rows[0][2] if s.pl_rows else (300, 200, 200)
        inch = s.erp_inches if s.erp_inches is not None else inches(true_box)
        # The ERP exports numbers as text; the reader must cope.
        ws.append([s.sku, s.name, str(rng.choice([4, 6, 8, 10, 12])), str(s.erp_qty),
                   "Ignore", s.erp_text, *[str(v) for v in inch], "Poly bag", None])
    _save(wb, INPUT / "erp_product_export_2026-09-18.xlsx")


PROSE = [
    "PACKING INSTRUCTIONS - please read before packing",
    "1. Every bag must carry the marketplace label on the outside.",
    "2. Do not mix SKUs inside one carton.",
    "3. Carton weight must not exceed 22 kg.",
    "4. Print the marketplace SKU and carton count on two sides of each carton.",
    "5. Use the bags-per-carton and carton size below. Ask before changing either.",
    "6. Send photos of three packed cartons before loading.",
]


def write_instructions(po: int, date: str, attr: str, with_internal: bool) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Box Information"
    for line in PROSE + ([""] * (3 if po == 1042 else 0)):
        ws.append([line])
    if with_internal:
        ws.append(["No.", "Internal SKU", "Parent SKU", "Product", "Pcs per bag",
                   "Bags per carton", "LENGTH (mm)", "WIDTH (mm)", "HEIGHT (mm)"])
    else:
        ws.append(["No.", "Parent SKU", "Marketplace SKU", "Product", "Pcs per bag",
                   "Bags (Boxes) per Carton", "L (Inch)", "W (Inch)", "H (Inch)",
                   "LENGTH (cm)", "WIDTH (cm)", "HEIGHT (cm)"])
    n = 0
    for s in SKUS:
        spec = getattr(s, attr)
        if spec is None:
            continue
        n += 1
        bags, box = spec
        if with_internal:
            ws.append([n, s.sku, s.parent, s.name, s.pcs_per_bag, bags, *(box or ("",) * 3)])
        else:
            # Header says (cm), values are millimetres; inch columns are the check.
            ws.append([n, s.parent, MKT[s.sku], s.name, s.pcs_per_bag, bags,
                       *(inches(box) if box else ("",) * 3), *(box or ("",) * 3)])
    if not with_internal:
        # A marketplace SKU with a typo: cannot be translated, must be flagged.
        ws.append([n + 1, "TW-5150", "ZZ9-TYP-0X0", "Twist Ties 50pc", 50, 100, "", "", "",
                   "", "", ""])
    ws.append(["TOTAL", None, None, None, None, None])
    _save(wb, INPUT / f"Packing Instructions Northwind PO {po} {date}.xlsx")


def write_packing_list() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PACKING LIST"
    ws.append(["NORTHWIND MANUFACTURING CO., LTD."])
    ws.append(["PACKING LIST - PO 1041 / PO 1042"])
    ws.append([])
    # The internal-SKU column has NO header; "FBA SKU" holds marketplace SKUs.
    ws.append(["No.", "ITEM NO.", "FBA SKU", None, "DESCRIPTION", "PCS/BAG", "BAGS/CTN",
               "CTNS", "QUANTITY", "LENGTH (cm)", "WIDTH (cm)", "HEIGHT (cm)", "MEAS (M3)",
               "G.W. (KG)"])
    n, total_ctns, total_qty = 0, 0, 0
    for s in SKUS:
        first_row = ws.max_row + 1
        for bags, ctns, box in s.pl_rows:
            n += 1
            qty = s.pcs_per_bag * bags * ctns + (5 if s.qty_typo else 0)
            cbm = round(box[0] * box[1] * box[2] / 1e9 * ctns, 3)
            internal = "#N/A" if s.na_in_pl else s.sku
            ws.append([n, s.parent, MKT[s.sku], internal, s.name, s.pcs_per_bag, bags, ctns,
                       qty, *box, cbm, round(rng.uniform(8, 18) * ctns, 1)])
            total_ctns += ctns
            total_qty += qty
        if len(s.pl_rows) > 1:
            # Same SKU in two packings: vendor merged the identity cells.
            for col in "BCDE":
                ws.merge_cells(f"{col}{first_row}:{col}{ws.max_row}")
    ws.append(["TOTAL", None, None, None, None, None, None, total_ctns, total_qty])
    _save(wb, INPUT / "Northwind packing list PO 1041-1042 2026-09-20.xlsx")


def write_decoys() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Qty"
    ws.append(["Product ID", "Product name", "QtyPerCase (ERP)", "Correct case qty", "Status"])
    ws.append(["FBA-HX-1042-P5", "Hex Key Set 10pc, 5-pack", 40, 40, "OK"])
    _save(wb, INPUT / "PackingCaseQty Northwind PO 1041 2026-09-01.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Unpacking"
    ws.append(["SKU", "PCS/BAG", "CTNS", "BAG QTY", "QUANTITY", "LOCATION"])
    ws.append(["HX-1042", 5, 12, 480, 2400, "Bay 3"])
    _save(wb, INPUT / "Unpacking list Northwind PO 1041-1042.xlsx")


def main() -> None:
    for folder in (INPUT, REFERENCE):
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*.xlsx"):
            old.unlink()
    write_catalog()
    write_erp_export()
    write_instructions(1041, "2026-08-14", "instr_1041", with_internal=True)
    write_instructions(1042, "2026-09-02", "instr_1042", with_internal=False)
    write_packing_list()
    write_decoys()


if __name__ == "__main__":
    main()
